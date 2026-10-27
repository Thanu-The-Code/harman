"""Read .xlsx/.xlsm (openpyxl) and .xls (xlrd) into one neutral Sheet model.

Only what form interpretation needs: displayed cell text (number formats applied),
merged ranges, which cell edges have borders, fonts, sizes, hidden rows/cols, print area,
embedded pictures, and the signs that mark a cell as a place to fill in: a background fill,
an unlocked cell, data validation (drop-down lists, dates, numbers), the number format of an
empty cell, and Form Control check boxes. Formulas are kept so they are never offered as inputs.
"""
import io
import re
import warnings
import zipfile
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from datetime import date, datetime, time

from ...errors import UnsupportedFile
from .package import Package

DEFAULT_COL_WIDTH = 8.43   # characters
DEFAULT_ROW_HEIGHT = 15.0  # points
MAX_OPTIONS = 200          # options read for one drop-down list
# Limits against hostile or runaway files (the upload limit applies to the compressed size only).
MAX_UNZIPPED = 300 * 1024 * 1024  # bytes inside an .xlsx package
MAX_RATIO = 200                   # uncompressed / compressed, per part ("zip bomb")
MAX_CELLS = 1_000_000             # stored cells across the workbook
WHITE = {"FFFFFF"}
WHITE_INDEXES = {1, 9, 64, 65}  # white / system background / system foreground in the legacy palette


@dataclass
class Validation:
    kind: str                     # "list", "date", "whole", "decimal", "time", "textLength", "custom"
    ranges: list                  # [(r0, r1, c0, c1)] end-exclusive, 0-based; r1/c1 may be None (whole column/row)
    options: list = None          # drop-down entries for "list"

    def covers(self, r, c):
        return any(r0 <= r and (r1 is None or r < r1) and c0 <= c and (c1 is None or c < c1) for r0, r1, c0, c1 in self.ranges)


@dataclass
class Sheet:
    name: str
    n_rows: int
    n_cols: int
    values: dict = field(default_factory=dict)       # (r, c) -> raw value
    texts: dict = field(default_factory=dict)        # (r, c) -> displayed text
    borders: dict = field(default_factory=dict)      # (r, c) -> set of "l","r","t","b"
    fonts: dict = field(default_factory=dict)        # (r, c) -> (size_pt, bold)
    merged: list = field(default_factory=list)       # (r0, r1, c0, c1) end-exclusive
    col_widths: dict = field(default_factory=dict)   # c -> characters
    row_heights: dict = field(default_factory=dict)  # r -> points
    hidden_rows: set = field(default_factory=set)
    hidden_cols: set = field(default_factory=set)
    print_area: tuple = None                         # (r0, r1, c0, c1) end-exclusive
    images: list = field(default_factory=list)       # {"row","col","src","width","height"}
    default_col_width: float = DEFAULT_COL_WIDTH
    default_row_height: float = DEFAULT_ROW_HEIGHT
    fills: dict = field(default_factory=dict)        # (r, c) -> background colour key (visible fills only)
    unlocked: set = field(default_factory=set)       # (r, c) explicitly unlocked for editing
    number_formats: dict = field(default_factory=dict)  # (r, c) -> format of cells that have no value yet
    formulas: dict = field(default_factory=dict)     # (r, c) -> "=SUM(B2:B5)"
    validations: list = field(default_factory=list)  # [Validation]
    controls: list = field(default_factory=list)     # Form Control check boxes / option buttons

    def validation_at(self, r, c):
        return next((dv for dv in self.validations if dv.covers(r, c)), None)


# ---- value formatting ------------------------------------------------------------------

_DATE_TOKENS = [("yyyy", "%Y"), ("yy", "%y"), ("mmmm", "%B"), ("mmm", "%b"), ("dd", "%d"), ("d", "%-d"), ("mm", "%m"), ("m", "%-m")]


def _date_text(value, fmt):
    f = (fmt or "").lower().split(";")[0]
    f = re.sub(r"\[[^\]]*\]|\"[^\"]*\"|\\\\.", "", f)
    if not any(t in f for t in ("d", "y")) or "h" in f and "d" not in f:
        f = "dd-mm-yyyy"
    out, i = "", 0
    while i < len(f):
        for tok, code in _DATE_TOKENS:
            if f.startswith(tok, i):
                if code == "%-d":
                    out += str(value.day)
                elif code == "%-m":
                    out += str(value.month)
                else:
                    out += value.strftime(code)
                i += len(tok)
                break
        else:
            out += f[i] if f[i] not in "hs" else ""
            i += 1
    return out.strip(" :")


def display_text(value, fmt="General"):
    """Text as Excel shows it (percentages, decimals, dates)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return _date_text(value, fmt)
    if isinstance(value, date):
        return _date_text(datetime(value.year, value.month, value.day), fmt)
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):
        f = (fmt or "General").split(";")[0]
        decimals = len(m.group(1)) if (m := re.search(r"0\.(0+)", f)) else 0
        if "%" in f:
            return f"{value * 100:.{decimals}f}%"
        if decimals:
            return f"{value:.{decimals}f}"
        if float(value).is_integer():
            return str(int(value))
        return f"{value:.10g}"
    return str(value)


def _ref(r, c):
    s, n = "", c + 1
    while n:
        n, rem = divmod(n - 1, 26)
        s = chr(65 + rem) + s
    return f"{s}{r + 1}"


def cell_ref(r0, c0, r1=None, c1=None):
    """0-based end-exclusive range -> 'A1' or 'A1:C3'."""
    if r1 is None or (r1 - r0 <= 1 and c1 - c0 <= 1):
        return _ref(r0, c0)
    return f"{_ref(r0, c0)}:{_ref(r1 - 1, c1 - 1)}"


def parse_ref(ref):
    """'F9', '$F$9' or 'F9:H12' -> (row, col) of the top-left cell, 0-based."""
    head = ref.split("!")[-1].split(":")[0].replace("$", "")
    letters = "".join(ch for ch in head if ch.isalpha())
    digits = "".join(ch for ch in head if ch.isdigit())
    col = 0
    for ch in letters.upper():
        col = col * 26 + (ord(ch) - 64)
    return int(digits) - 1, col - 1


def parse_sqref(sqref):
    """'B4 D4:D10 F:F' -> [(r0, r1, c0, c1)] end-exclusive; None for an open end."""
    from openpyxl.utils.cell import range_boundaries

    out = []
    for part in (sqref or "").split():
        try:
            c0, r0, c1, r1 = range_boundaries(part.replace("$", ""))
        except ValueError:
            continue
        out.append(((r0 or 1) - 1, r1, (c0 or 1) - 1, c1))
    return out


def is_number_format(fmt):
    f = (fmt or "").split(";")[0]
    return bool(re.search(r"[0#]", re.sub(r"\"[^\"]*\"|\[[^\]]*\]", "", f)))


# ---- readers -----------------------------------------------------------------------------

class WorkbookReader(ABC):
    @abstractmethod
    def read(self, data: bytes) -> list:
        """-> [Sheet] for visible worksheets"""


def _xlsx_fill_key(fill):
    """Colour key of a visible background fill, or None for no fill / white."""
    if fill is None or not fill.fill_type or fill.fill_type == "none":
        return None
    if fill.fill_type != "solid":
        return f"pattern:{fill.fill_type}"
    col = fill.fgColor
    if col is None:
        return None
    if col.type == "rgb":
        rgb = (col.rgb or "")[-6:].upper() if isinstance(col.rgb, str) else ""
        return None if not rgb or rgb in WHITE else rgb
    if col.type == "theme":
        if col.theme == 0 and not col.tint:
            return None  # theme background (white)
        return f"theme:{col.theme}:{round(col.tint or 0, 2)}"
    if col.type == "indexed":
        return None if col.indexed in WHITE_INDEXES else f"indexed:{col.indexed}"
    return None


class XlsxReader(WorkbookReader):
    def _load(self, data, data_only):
        import openpyxl

        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)

    def read(self, data):
        from openpyxl.utils.cell import range_boundaries

        check_package(data)
        try:
            zf = zipfile.ZipFile(io.BytesIO(data))
            pkg = Package(zf)
            wb = self._load(data, data_only=True)
        except Exception as exc:  # zipfile / XML errors
            raise UnsupportedFile(f"Could not read Excel file: {exc}") from exc
        pictures, controls, x14 = pkg.pictures(), pkg.controls(), pkg.x14_validations()
        has_formulas = any(re.search(rb"<f[ >]", zf.read(part)) for part in pkg.sheet_parts.values())
        formula_wb = self._load(data, data_only=False) if has_formulas else None
        check_cells(sum(len(ws._cells) for ws in wb.worksheets), "The workbook")  # noqa: SLF001
        sheets = []
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            sh = Sheet(name=ws.title, n_rows=ws.max_row, n_cols=ws.max_column)
            fmt = ws.sheet_format
            sh.default_row_height = float(fmt.defaultRowHeight or DEFAULT_ROW_HEIGHT)
            sh.default_col_width = float(fmt.defaultColWidth or ((fmt.baseColWidth or 8) + 0.43))
            # ws._cells holds only the cells stored in the file; iter_rows() would create every
            # cell of the max_row x max_column rectangle.
            for (r1, c1), cell in ws._cells.items():  # noqa: SLF001
                r, c = r1 - 1, c1 - 1
                if cell.value is not None:
                    sh.values[(r, c)] = cell.value
                    sh.texts[(r, c)] = display_text(cell.value, cell.number_format)
                    sh.fonts[(r, c)] = (float(cell.font.sz or 11), bool(cell.font.b))
                if not cell.has_style:
                    continue
                b = cell.border
                sides = {k for k, side in (("l", b.left), ("r", b.right), ("t", b.top), ("b", b.bottom)) if side is not None and side.style}
                if sides:
                    sh.borders[(r, c)] = sides
                key = _xlsx_fill_key(cell.fill)
                if key:
                    sh.fills[(r, c)] = key
                if cell.protection is not None and cell.protection.locked is False:
                    sh.unlocked.add((r, c))
                if cell.value is None and cell.number_format not in ("General", "@"):
                    sh.number_formats[(r, c)] = cell.number_format
            if formula_wb is not None and ws.title in formula_wb.sheetnames:
                for (r1, c1), cell in formula_wb[ws.title]._cells.items():  # noqa: SLF001
                    if cell.data_type == "f" and isinstance(cell.value, str):
                        sh.formulas[(r1 - 1, c1 - 1)] = cell.value
            for rng in ws.merged_cells.ranges:
                sh.merged.append((rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col))
            for dim in ws.column_dimensions.values():
                lo, hi = (dim.min or 1) - 1, (dim.max or dim.min or 1)
                for c in range(lo, min(hi, lo + 16384)):
                    if dim.width:
                        sh.col_widths[c] = float(dim.width)
                    if dim.hidden:
                        sh.hidden_cols.add(c)
            for r, dim in ws.row_dimensions.items():
                if dim.height is not None:
                    sh.row_heights[r - 1] = float(dim.height)
                if dim.hidden:
                    sh.hidden_rows.add(r - 1)
            area = ws.print_area
            if area:
                boxes = []
                for piece in (area if isinstance(area, (list, tuple)) else area.split(",")):
                    c0, r0, c1, r1 = range_boundaries(piece.split("!")[-1].replace("$", ""))
                    boxes.append((r0 - 1, r1, c0 - 1, c1))
                # Several print areas: the form is the rectangle that holds them all.
                sh.print_area = (min(b[0] for b in boxes), max(b[1] for b in boxes),
                                 min(b[2] for b in boxes), max(b[3] for b in boxes))
            for dv in ws.data_validations.dataValidation:
                self._add_validation(sh, wb, ws, dv.type, dv.formula1, str(dv.sqref))
            for kind, formula1, sqref in x14.get(ws.title, []):
                self._add_validation(sh, wb, ws, kind, formula1, sqref)
            sh.images = pictures.get(ws.title, [])
            sh.controls = controls.get(ws.title, [])
            sheets.append(sh)
        return sheets

    def _add_validation(self, sh, wb, ws, kind, formula1, sqref):
        if not kind or kind == "none":
            return
        ranges = parse_sqref(sqref)
        if not ranges:
            return
        options = self._list_options(wb, ws, formula1) if kind == "list" else None
        sh.validations.append(Validation(kind=kind, ranges=ranges, options=options))

    @staticmethod
    def _list_options(wb, ws, formula):
        """Entries of a drop-down list: inline ("Yes,No"), a range (also on a hidden sheet), or a named range."""
        from openpyxl.utils.cell import range_boundaries

        f = (formula or "").strip()
        if len(f) >= 2 and f[0] == '"' and f[-1] == '"':
            return [x.strip() for x in f[1:-1].split(",") if x.strip()][:MAX_OPTIONS]
        f = f.lstrip("=")
        named = ws.defined_names.get(f) if hasattr(ws, "defined_names") else None
        named = named or wb.defined_names.get(f)
        if named is not None:
            f = (named.attr_text or "").lstrip("=")
        m = re.match(r"^(?:(?P<sheet>'[^']+'|[^'!]+)!)?(?P<range>\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?)$", f)
        if not m:
            return None  # INDIRECT(), OFFSET() and other computed lists
        sheet_name = (m.group("sheet") or "").strip("'")
        target = wb[sheet_name] if sheet_name in wb.sheetnames else ws
        try:
            c0, r0, c1, r1 = range_boundaries(m.group("range").replace("$", ""))
        except ValueError:
            return None
        out = []
        for row in target.iter_rows(min_row=r0, max_row=min(r1, r0 + MAX_OPTIONS), min_col=c0, max_col=min(c1, c0 + 20)):
            for cell in row:
                text = display_text(cell.value, cell.number_format).strip()
                if text and text not in out:
                    out.append(text)
        return out[:MAX_OPTIONS]


class XlsReader(WorkbookReader):
    @staticmethod
    def _fill_key(bk, xf):
        bg = xf.background
        if not bg or not bg.fill_pattern:
            return None
        idx = bg.pattern_colour_index if bg.fill_pattern == 1 else None
        if idx is None:
            return f"pattern:{bg.fill_pattern}"
        rgb = bk.colour_map.get(idx)
        if idx in WHITE_INDEXES or rgb is None or tuple(rgb) == (255, 255, 255):
            return None
        return "%02X%02X%02X" % tuple(rgb)

    def read(self, data):
        import xlrd

        try:
            bk = xlrd.open_workbook(file_contents=data, formatting_info=True)
        except Exception as exc:
            raise UnsupportedFile(f"Could not read Excel file: {exc}") from exc
        areas = {}
        for name in bk.name_obj_list:
            if name.name.lower() == "print_area":
                try:
                    sheet, r0, r1, c0, c1 = name.area2d()
                    areas[sheet.name] = (r0, r1, c0, c1)
                except Exception:  # noqa: BLE001 - complex/multi-area names are ignored
                    pass
        check_cells(sum(s.nrows * s.ncols for s in bk.sheets()), "The workbook")
        sheets = []
        for sht in bk.sheets():
            if sht.visibility != 0:
                continue
            sh = Sheet(name=sht.name, n_rows=sht.nrows, n_cols=sht.ncols)
            sh.default_row_height = (sht.default_row_height or 300) / 20.0
            sh.default_col_width = float(sht.defcolwidth or 8) + 0.43
            for r in range(sht.nrows):
                for c in range(sht.ncols):
                    cell = sht.cell(r, c)
                    xf = bk.xf_list[sht.cell_xf_index(r, c)]
                    fmt = bk.format_map[xf.format_key].format_str if xf.format_key in bk.format_map else "General"
                    if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, bk.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = ""
                        sh.values[(r, c)] = value
                        sh.texts[(r, c)] = display_text(value, fmt)
                        font = bk.font_list[xf.font_index]
                        sh.fonts[(r, c)] = (font.height / 20.0, bool(font.bold))
                    elif fmt not in ("General", "@"):
                        sh.number_formats[(r, c)] = fmt
                    bd = xf.border
                    sides = {k for k, style in (("l", bd.left_line_style), ("r", bd.right_line_style),
                                                ("t", bd.top_line_style), ("b", bd.bottom_line_style)) if style}
                    if sides:
                        sh.borders[(r, c)] = sides
                    key = self._fill_key(bk, xf)
                    if key:
                        sh.fills[(r, c)] = key
                    if xf.protection is not None and not xf.protection.cell_locked:
                        sh.unlocked.add((r, c))
            sh.merged = [(r0, r1, c0, c1) for r0, r1, c0, c1 in sht.merged_cells]
            for c, info in sht.colinfo_map.items():
                sh.col_widths[c] = info.width / 256.0
                if info.hidden:
                    sh.hidden_cols.add(c)
            for r, info in sht.rowinfo_map.items():
                sh.row_heights[r] = info.height / 20.0
                if info.hidden:
                    sh.hidden_rows.add(r)
            sh.print_area = areas.get(sht.name)
            sheets.append(sh)
        return sheets


def is_xlsx(data):
    if data[:2] != b"PK":
        return False
    try:
        return "xl/workbook.xml" in zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile:
        return False


def check_package(data):
    """Refuse packages that would expand to far more than any form needs."""
    try:
        infos = zipfile.ZipFile(io.BytesIO(data)).infolist()
    except zipfile.BadZipFile as exc:
        raise UnsupportedFile(f"Could not read Excel file: {exc}") from exc
    total = sum(i.file_size for i in infos)
    if total > MAX_UNZIPPED:
        raise UnsupportedFile(f"The workbook expands to {total // (1024 * 1024)} MB; the limit is {MAX_UNZIPPED // (1024 * 1024)} MB")
    for i in infos:
        if i.compress_size and i.file_size > 1024 * 1024 and i.file_size / i.compress_size > MAX_RATIO:
            raise UnsupportedFile("The workbook looks malformed (a part expands far beyond its size)")


def check_cells(count, where):
    if count > MAX_CELLS:
        raise UnsupportedFile(f"{where} has {count:,} cells; forms up to {MAX_CELLS:,} cells are supported")


def is_xls(data):
    return data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def read_workbook(data):
    if is_xlsx(data):
        sheets = XlsxReader().read(data)
    elif is_xls(data):
        sheets = XlsReader().read(data)
    else:
        raise UnsupportedFile("Not an Excel workbook (.xlsx, .xlsm or .xls)")
    if not sheets:
        raise UnsupportedFile("The workbook has no visible sheets")
    return sheets
