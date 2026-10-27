"""Writing a form's values back into the workbook (POST /ocr/export)."""
import base64
import io
import json
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.ocr.formats import UploadedFile, build_registry
from app.ocr.formats.excel.workbook import read_workbook
from app.ocr.values import plain

from . import client_samples as cs
from .excel_forms import FILLED, employee_form


@pytest.fixture(scope="module")
def handler():
    return build_registry("excel").get("excel")


@pytest.fixture(scope="module")
def employee(handler):
    blank = employee_form()
    td = handler.build_template(UploadedFile(blank, "form.xlsx"))["task_data"]
    answers = handler.extract(UploadedFile(employee_form(FILLED), "f.xlsx"), td)["answers"]
    out = handler.export(UploadedFile(blank, "onboarding.xlsx"), td, answers)
    return blank, td, answers, out


def test_round_trip(handler, employee):
    _, td, answers, out = employee
    assert out["filename"] == "onboarding-filled.xlsx" and not out["skipped"]
    back = handler.extract(UploadedFile(out["data"], "x.xlsx"), td)["answers"]
    assert back == answers


def test_values_are_typed_cells(employee):
    ws = openpyxl.load_workbook(io.BytesIO(employee[3]["data"]))["Form"]
    assert ws["B7"].value.year == 2025 and ws["B7"].number_format == "dd/mm/yyyy"  # a real date
    assert ws["B8"].value == 29  # a number
    assert ws["B4"].value == "Female"  # the drop-down entry, not its slug
    assert ws["A19"].value == "Married: ☐ Yes ☑ No" and ws["C21"].value == "☑"
    assert ws["B11"].value == "=B8*2"  # formulas are kept, never overwritten


def test_only_touched_parts_change(employee):
    blank, _, _, out = employee
    src, dst = zipfile.ZipFile(io.BytesIO(blank)), zipfile.ZipFile(io.BytesIO(out["data"]))
    assert src.namelist() == dst.namelist()
    changed = {n for n in src.namelist() if src.read(n) != dst.read(n)}
    assert changed <= {"xl/worksheets/sheet1.xml", "xl/workbook.xml", "xl/drawings/vmlDrawing1.vml",
                       "xl/ctrlProps/ctrlProp1.xml", "xl/ctrlProps/ctrlProp2.xml", "xl/ctrlProps/ctrlProp3.xml",
                       "xl/ctrlProps/ctrlProp4.xml"}, changed
    for name in changed:
        ET.fromstring(dst.read(name))  # still well-formed XML
    assert b'fullCalcOnLoad="1"' in dst.read("xl/workbook.xml")  # formulas recalculate on open


def test_form_controls_are_ticked(employee):
    sheet = next(s for s in read_workbook(employee[3]["data"]) if s.name == "Form")
    state = {c["text"]: c["checked"] for c in sheet.controls}
    assert state == {"ID proof": True, "Address proof": True, "Yes": False, "No": True}
    assert sheet.values[(23, 7)] is True  # the ID proof control's linked cell H24


def test_unticking_clears_boxes(handler, employee):
    _, td, answers, out = employee
    by = {plain(e.get("label", "")): e for e in td if e.get("field_name")}
    cleared = {**answers, by["Married"]["field_name"]: [], by["Documents received"]["field_name"]: []}
    again = handler.export(UploadedFile(out["data"], "x.xlsx"), td, cleared)
    sheet = next(s for s in read_workbook(again["data"]) if s.name == "Form")
    assert sheet.texts[(18, 0)] == "Married: ☐ Yes ☐ No"
    assert {c["text"]: c["checked"] for c in sheet.controls}["ID proof"] is False


def test_template_without_cell_addresses(handler, employee):
    """A form drawn in the builder (or made from a scan) is written by finding labels/headings."""
    blank, td, answers, _ = employee
    bare = json.loads(json.dumps(td))
    for e in bare:
        e.pop("excel", None)
    out = handler.export(UploadedFile(blank, "form.xlsx"), bare, answers)
    back = handler.extract(UploadedFile(out["data"], "x.xlsx"), td)["answers"]
    labels = {e["field_name"]: plain(e.get("label", "")) for e in td if e.get("field_name")}
    for name in ("Name", "Age", "Email", "Married", "Shift", "Documents received", "Own vehicle"):
        key = next(k for k, label in labels.items() if label == name)
        assert back.get(key) == answers[key], name


@pytest.mark.skipif(not cs.available(), reason="client samples (tests/ocr/samples/client) not present")
@pytest.mark.parametrize("key", list(cs.SAMPLES))
def test_client_workbooks_round_trip(handler, key):
    ext = cs.EXCEL_EXT[key]
    data = cs.read(key, ext)
    template = handler.build_template(UploadedFile(data, "blank" + ext))
    td = template["task_data"]
    want = cs.expected_answers(td, cs.plan(td))
    out = handler.export(UploadedFile(data, "blank" + ext), td, want)
    assert out["filename"].endswith("-filled" + ext) and not out["skipped"]
    got = handler.extract(UploadedFile(out["data"], "x" + ext), td)["answers"]
    assert {k: got.get(k) for k in want} == want
    # Formatting (borders, merges) is kept: the filled workbook still reads as the same form.
    again = handler.read_form(UploadedFile(out["data"], "x" + ext))["task_data"]
    shape = lambda t: sorted((e["element"], len(e.get("columns") or []), len(e.get("rows") or [])) for e in t if e["element"] == "Table")  # noqa: E731
    assert shape(again) == shape(td)
    assert len([e for e in again if e["element"] == "Image"]) == len([e for e in td if e["element"] == "Image"])


def test_export_endpoint():
    from .api_app import app

    client = TestClient(app)
    blank = employee_form()
    td = client.post("/ocr/template", files={"file": ("form.xlsx", blank)}).json()["task_data"]
    name = next(e["field_name"] for e in td if plain(e.get("label", "")) == "Name")
    r = client.post("/ocr/export", files={"file": ("form.xlsx", blank)},
                    data={"form": json.dumps(td), "answers": json.dumps([{"name": name, "value": "Ravi"}])})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["filename"] == "form-filled.xlsx" and body["written"] == ["Name"]
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(body["data"])))["Form"]
    assert ws["B3"].value == "Ravi"
    formats = {f["key"]: f["capabilities"] for f in client.get("/ocr/formats").json()["formats"]}
    assert "export" in formats["excel"] and "export" not in formats.get("pdf", [])
    pdf = client.post("/ocr/export", files={"file": ("x.pdf", b"%PDF-1.4 x")}, data={"form": "[]", "answers": "{}"})
    assert pdf.status_code == 415
