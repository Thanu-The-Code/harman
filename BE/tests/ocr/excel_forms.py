"""Synthetic Excel forms that use the ways real workbooks mark inputs (beyond borders):
data validation (inline, cross-sheet Excel 2010+, named range), number formats, shaded and
unlocked cells, typed tick boxes, Form Control check boxes / option buttons, and formulas.

openpyxl cannot write Form Controls or Excel 2010+ validation, so those parts are injected into
the saved package the way Excel stores them (VML drawing, ctrlProps, x14 extLst).
"""
import io
import re
import zipfile
from datetime import date

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SHADE = PatternFill("solid", fgColor="FFF2CC")
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

# Form Controls: (caption, kind, anchor cell (row, col) 0-based, linked cell)
CONTROLS = [
    ("ID proof", "Checkbox", (23, 2), "$H$24"),
    ("Address proof", "Checkbox", (23, 4), None),
    ("Yes", "Radio", (24, 2), None),
    ("No", "Radio", (24, 4), None),
]

BLANK = {
    "married": "Married: ☐ Yes ☐ No",
    "smoker": ("☐", "☐"),
    "vehicle": ("☐", "☐"),
    "shift": "Shift: ○ Day ○ Night",
    "controls": [False, False, False, False],
    "ctrlprops": [None, None, None, None],
    "links": {},
}
FILLED = {
    "B3": "Asha Rao", "B4": "Female", "B5": "Quality", "B6": "G2", "B7": date(2025, 3, 14), "B8": 29,
    "B9": date(2024, 11, 2), "B10": 9876543210, "B13": "98450 12345", "B14": "asha@example.com",
    "B15": "12 MG Road, Pune",
    "C28": "✓", "D28": date(2025, 1, 5), "C29": "✗",
    "married": "Married: ☐ Yes ☑ No",
    "smoker": ("☐", "☑"),
    "vehicle": ("☑", "☐"),
    "shift": "Shift: ○ Day ● Night",
    # ID proof ticked in the VML, Address proof only in ctrlProps (Excel 2010+), "No" radio ticked
    "controls": [True, False, False, True],
    "ctrlprops": [None, "Checked", None, None],
    "links": {"H24": True},
}


def employee_form(filled=None):
    """The onboarding form, blank or filled (`filled` maps cells / keys to values)."""
    state = {**BLANK, **(filled or {})}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Form"
    lists = wb.create_sheet("Lists")
    for i, dept in enumerate(["Quality", "Production", "Stores"], start=1):
        lists.cell(i, 1, dept)
    for i, grade in enumerate(["G1", "G2", "G3"], start=1):
        lists.cell(i, 2, grade)
    lists.sheet_state = "hidden"
    wb.defined_names["Grades"] = DefinedName("Grades", attr_text="Lists!$B$1:$B$3")
    for col, width in zip("ABCDEFGH", (22, 18, 12, 14, 12, 10, 6, 8)):
        ws.column_dimensions[col].width = width

    ws["A1"] = "EMPLOYEE ONBOARDING FORM"
    ws["A1"].font = Font(size=16, bold=True)
    ws.merge_cells("A1:F1")

    # Bordered key/value rows; the value cells say what they accept.
    rows = [("Name:", None), ("Gender:", "list"), ("Department:", "x14"), ("Grade:", "named"),
            ("Date of joining:", "date"), ("Age:", "whole"), ("Joined on", "datefmt"), ("Phone:", "numfmt"),
            ("Total", "formula")]
    for r, (label, kind) in enumerate(rows, start=3):
        ws.cell(r, 1, label).border = BOX
        c = ws.cell(r, 2)
        c.border = BOX
        if kind == "list":
            dv = DataValidation(type="list", formula1='"Male,Female,Other"')
            ws.add_data_validation(dv)
            dv.add(c)
        elif kind == "named":
            dv = DataValidation(type="list", formula1="Grades")
            ws.add_data_validation(dv)
            dv.add(c)
        elif kind == "date":
            dv = DataValidation(type="date", operator="greaterThan", formula1="1")
            ws.add_data_validation(dv)
            dv.add(c)
            c.number_format = "dd/mm/yyyy"
        elif kind == "whole":
            dv = DataValidation(type="whole", operator="between", formula1="18", formula2="70")
            ws.add_data_validation(dv)
            dv.add(c)
        elif kind == "datefmt":
            c.number_format = "dd-mmm-yy"
        elif kind == "numfmt":
            c.number_format = "0"
        elif kind == "formula":
            c.value = "=B8*2"
    # Shaded, unlocked inputs with no borders (sheet protected).
    for r, label in ((13, "Mobile"), (14, "Email"), (15, "Address")):
        ws.cell(r, 1, label)
        for rr in ((15, 16, 17) if label == "Address" else (r,)):
            for cc in (2, 3, 4):
                cell = ws.cell(rr, cc)
                cell.fill = SHADE
                cell.protection = Protection(locked=False)
    ws.protection.sheet = True

    # Typed tick boxes.
    ws["A19"] = state["married"]
    ws.merge_cells("A19:C19")
    ws["A20"], ws["C20"], ws["E20"] = "Smoker", "Yes", "No"
    ws["B20"], ws["D20"] = state["smoker"]
    ws["A21"], ws["B21"], ws["D21"] = "Own vehicle", "Yes", "No"
    ws["C21"], ws["E21"] = state["vehicle"]
    ws["A22"] = state["shift"]
    ws.merge_cells("A22:C22")
    # Form Controls go in rows 24-25 (injected below).
    ws["A24"] = "Documents received"
    ws["A25"] = "Laptop issued"
    for ref, value in state["links"].items():
        ws[ref] = value

    # A checklist table: a tick column (✓/✗ list) and a date column (date format).
    for c, text in enumerate(["Sr. No.", "Item", "Checked", "Date"], start=1):
        ws.cell(27, c, text).border = BOX
    tick = DataValidation(type="list", formula1='"✓,✗"')
    ws.add_data_validation(tick)
    for r, item in zip((28, 29, 30), ("Laptop", "ID card", "Locker key")):
        ws.cell(r, 1, r - 27).border = BOX
        ws.cell(r, 2, item).border = BOX
        for c in (3, 4):
            ws.cell(r, c).border = BOX
        tick.add(ws.cell(r, 3))
        ws.cell(r, 4).number_format = "dd/mm/yyyy"
    ws["A32"] = "Note: return the signed copy to HR."

    for ref, value in state.items():
        if re.fullmatch(r"[A-H]\d+", ref):
            ws[ref] = value
    bio = io.BytesIO()
    wb.save(bio)
    return _inject(bio.getvalue(), state)


def _vml(state):
    shapes = []
    for n, ((caption, kind, (r, c), link), checked) in enumerate(zip(CONTROLS, state["controls"])):
        extra = "<x:Checked>1</x:Checked>" if checked else ""
        extra += f"<x:FmlaLink>{link}</x:FmlaLink>" if link else ""
        shapes.append(f"""
 <v:shape id="Control_{n}" type="#_x0000_t201" style="position:absolute;z-index:{n + 1}" filled="f" stroked="f" o:insetmode="auto" o:spid="_x0000_s{1025 + n}">
  <v:path shadowok="t" strokeok="t" fillok="t"/>
  <v:textbox style="mso-direction-alt:auto" o:singleclick="f">
   <div style="text-align:left"><font face="Segoe UI" size="160" color="auto">{caption}</font></div>
  </v:textbox>
  <x:ClientData ObjectType="{kind}">
   <x:Anchor>{c}, 4, {r}, 2, {c + 1}, 30, {r}, 18</x:Anchor>
   <x:AutoFill>False</x:AutoFill>{extra}
  </x:ClientData>
 </v:shape>""")
    return ('<xml xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:x="urn:schemas-microsoft-com:office:excel">\n'
            ' <v:shapetype id="_x0000_t201" coordsize="21600,21600" o:spt="201" path="m,l,21600r21600,l21600,xe"/>'
            + "".join(shapes) + "\n</xml>")


def _inject(data, state):
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                body = body.replace(b"</Types>", b'<Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/></Types>')
            if item.filename == "xl/worksheets/sheet1.xml":
                controls = "".join(
                    f'<control xmlns:r="{R_NS}" shapeId="{1025 + n}" r:id="rIdC{n}" name="Control {n}"/>'
                    for n in range(len(CONTROLS)))
                ext = ('<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
                       'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                       '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
                       '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Lists!$A$1:$A$3</xm:f></x14:formula1>'
                       '<xm:sqref>B5</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
                body = body.replace(b"</worksheet>", (f'<legacyDrawing xmlns:r="{R_NS}" r:id="rIdV"/>'
                                                      f"<controls>{controls}</controls>{ext}</worksheet>").encode())
            dst.writestr(item, body)
        rels = ['<Relationship Id="rIdV" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" '
                'Target="../drawings/vmlDrawing1.vml"/>']
        for n, prop in enumerate(state["ctrlprops"]):
            rels.append(f'<Relationship Id="rIdC{n}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/ctrlProp" '
                        f'Target="../ctrlProps/ctrlProp{n + 1}.xml"/>')
            kind = "CheckBox" if CONTROLS[n][1] == "Checkbox" else "Radio"
            checked = f' checked="{prop}"' if prop else ""
            dst.writestr(f"xl/ctrlProps/ctrlProp{n + 1}.xml",
                         f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<formControlPr '
                         f'xmlns="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" objectType="{kind}"{checked} lockText="1"/>')
        dst.writestr("xl/worksheets/_rels/sheet1.xml.rels",
                     '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                     '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     + "".join(rels) + "</Relationships>")
        dst.writestr("xl/drawings/vmlDrawing1.vml", _vml(state))
    return out.getvalue()
