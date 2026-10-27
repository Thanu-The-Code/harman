"""Excel inputs marked by more than borders: data validation, number formats, shading, unlocked
cells, typed tick boxes, Form Controls and formulas (synthetic workbooks in excel_forms.py)."""
import copy
import io
import time

import openpyxl
import pytest
from openpyxl.styles import Border, PatternFill, Protection, Side

from app.ocr.formats import UploadedFile, build_registry
from app.ocr.values import plain

from .excel_forms import FILLED, employee_form

FIELDS = {
    "Name": ("TextInput", "B3"), "Gender": ("Dropdown", "B4"), "Department": ("Dropdown", "B5"),
    "Grade": ("Dropdown", "B6"), "Date of joining": ("DatePicker", "B7"), "Age": ("NumberInput", "B8"),
    "Joined on": ("DatePicker", "B9"), "Phone": ("PhoneNumber", "B10"), "Mobile": ("PhoneNumber", "B13:D13"),
    "Email": ("EmailInput", "B14:D14"), "Address": ("TextArea", "B15:D17"),
}
CHOICES = {
    "Married": ("Checkboxes", ["Yes", "No"]), "Smoker": ("Checkboxes", ["Yes", "No"]),
    "Own vehicle": ("Checkboxes", ["Yes", "No"]), "Shift": ("RadioButtons", ["Day", "Night"]),
    "Documents received": ("Checkboxes", ["ID proof", "Address proof"]), "Laptop issued": ("RadioButtons", ["Yes", "No"]),
}
VALUES = {
    "Name": "Asha Rao", "Gender": "Female", "Department": "Quality", "Grade": "G2", "Date of joining": "03/14/2025",
    "Age": "29", "Joined on": "11/02/2024", "Phone": "9876543210", "Mobile": "98450 12345", "Email": "asha@example.com",
    "Address": "12 MG Road, Pune", "Married": "No", "Smoker": "No", "Own vehicle": "Yes", "Shift": "Night",
    "Documents received": "ID proof, Address proof", "Laptop issued": "No",
}
TABLE = [["", "", "✓", "2025-01-05"], ["", "", "✗", ""], ["", "", "", ""]]


@pytest.fixture(scope="module")
def handler():
    return build_registry("excel").get("excel")


@pytest.fixture(scope="module")
def blank(handler):
    return handler.build_template(UploadedFile(employee_form(), "form.xlsx"))


def labelled(task_data):
    return {plain(e.get("label", "")): e for e in task_data if e.get("field_name") and e.get("label")}


def check_form(task_data):
    by = labelled(task_data)
    for label, (kind, ref) in FIELDS.items():
        assert label in by, f"missing {label}"
        assert (by[label]["element"], by[label]["excel"]["cell"]) == (kind, ref), label
    for label, (kind, options) in CHOICES.items():
        assert by[label]["element"] == kind, label
        assert [o["text"] for o in by[label]["options"]] == options, label
    assert [o["text"] for o in by["Gender"]["options"]] == ["Male", "Female", "Other"]
    assert [o["text"] for o in by["Department"]["options"]] == ["Quality", "Production", "Stores"]  # hidden sheet, x14
    assert [o["text"] for o in by["Grade"]["options"]] == ["G1", "G2", "G3"]  # named range
    assert "Total" not in by, "a formula is calculated, not filled in"
    table = next(e for e in task_data if e["element"] == "Table")
    assert [[c.get("input") for c in r["cells"]][2:] for r in table["rows"]] == [["check", "date"]] * 3
    assert "TRUE" not in [plain(e.get("content", "")) for e in task_data]  # a control's linked cell


def check_values(result):
    got = {f["label"]: f["value"] for f in result["fields"]}
    for label, want in VALUES.items():
        value = got.get(label)
        if isinstance(value, str) and label in ("Gender", "Department", "Grade"):
            assert value.lower() == want.lower(), label
        else:
            assert value == want, (label, value, want)
    table = next(f for f in result["fields"] if f["element"] == "Table")
    assert table["value"] == TABLE


def test_template_reads_input_markers(blank):
    check_form(blank["task_data"])


def test_fill_by_cell(handler, blank):
    result = handler.extract(UploadedFile(employee_form(FILLED), "filled.xlsx"), blank["task_data"])
    check_values(result)
    assert {f["method"] for f in result["fields"] if f["value"] is not None} == {"cell"}


def test_fill_by_labels_only(handler, blank):
    bare = copy.deepcopy(blank["task_data"])
    for e in bare:
        e.pop("excel", None)
    result = handler.extract(UploadedFile(employee_form(FILLED), "filled.xlsx"), bare)
    got = {f["label"]: f["value"] for f in result["fields"]}
    for label in ("Name", "Age", "Mobile", "Email", "Married", "Shift", "Documents received", "Laptop issued", "Own vehicle"):
        assert got.get(label) == VALUES[label], (label, got.get(label))


def test_filled_copy_reads_back_as_the_same_form(handler):
    result = handler.read_form(UploadedFile(employee_form(FILLED), "filled.xlsx"))
    check_form(result["task_data"])
    check_values(result)


def test_digital_copy_shows_tick_boxes(blank):
    page = blank["document"]["pages"][0]
    assert len(page["checkboxes"]) >= 6  # lone typed boxes and Form Controls
    assert any(line["text"] == "ID proof" for line in page["lines"])


# ---- guards: styling that does not mark inputs ---------------------------------------------------

def _sheet(build):
    wb = openpyxl.Workbook()
    ws = wb.active
    build(ws)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _elements(handler, data):
    return [(e["element"], plain(e.get("label") or e.get("content") or "")) for e in
            handler.build_template(UploadedFile(data, "x.xlsx"))["task_data"] if not e["element"].endswith("Row")]


def test_sheet_background_colour_is_not_an_input(handler):
    def build(ws):
        grey = PatternFill("solid", fgColor="EEEEEE")
        for r in range(1, 12):
            for c in range(1, 6):
                ws.cell(r, c).fill = grey
        ws["A1"], ws["A3"], ws["A4"], ws["C6"] = "Visitor log", "Name", "Company", "Thank you"
        thin = Side(style="thin")
        ws["B3"].border = Border(bottom=thin)
    elements = _elements(handler, _sheet(build))
    inputs = [x for x in elements if x[0] not in ("Header", "Label", "Paragraph")]
    assert inputs == [("TextInput", "Name")], elements  # only the underlined box


def test_unlocking_everything_does_not_mark_inputs(handler):
    def build(ws):
        for r in range(1, 10):
            for c in range(1, 5):
                ws.cell(r, c).protection = Protection(locked=False)
        ws["A1"], ws["A3"], ws["A4"] = "Notes", "Topic", "Owner"
    elements = _elements(handler, _sheet(build))
    assert not [x for x in elements if x[0] not in ("Header", "Label", "Paragraph")], elements


def test_thin_coloured_strip_is_decoration(handler):
    def build(ws):
        ws["A1"] = "Request"
        ws["A2"] = "Section A"
        ws.row_dimensions[3].height = 4
        for c in range(1, 6):
            ws.cell(3, c).fill = PatternFill("solid", fgColor="1F4E79")
    elements = _elements(handler, _sheet(build))
    assert not [x for x in elements if x[0] not in ("Header", "Label", "Paragraph")], elements


def test_formatting_dragged_far_down_stays_fast(handler):
    def build(ws):
        ws["A1"], ws["A3"] = "Short form", "Name:"
        for c in ("A3", "B3"):
            ws[c].border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        for r in range(1, 30001):
            ws.cell(r, 6).fill = PatternFill("solid", fgColor="FFFF99")
    data = _sheet(build)
    start = time.time()
    result = handler.build_template(UploadedFile(data, "x.xlsx"))
    assert time.time() - start < 3
    assert result["document"]["pages"][0]["height"] < 400  # the page is the form, not 30,000 rows
